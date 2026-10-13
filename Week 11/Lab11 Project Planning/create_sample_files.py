# agentic-document-automator/create_sample_files.py
# Creates the input files the agent needs:
#   templates/report_template.docx  - Word template with {{PLACEHOLDERS}}
#   documents/source_document.pdf   - PDF to watermark and extract text from
#   documents/dummy_doc1.pdf        - PDF part 1 for merging
#   documents/dummy_doc2.pdf        - PDF part 2 for merging
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.join(BASE_DIR, 'documents')
TEMPLATES_DIR = os.path.join(BASE_DIR, 'templates')


# ---------- Word template ----------
def create_word_template(path):
    doc = docx.Document()

    title = doc.add_heading('{{REPORT_TITLE}}', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run('Date: {{REPORT_DATE}}   |   Author: {{AUTHOR_NAME}}   |   {{COMPANY_NAME}}')
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph('{{INTRODUCTION_TEXT}}')

    doc.add_heading('2. Progress Overview: {{MONTH_PERIOD}}', level=1)
    doc.add_paragraph('The table below summarizes the status of each feature:')
    doc.add_paragraph('{{TABLE_PLACEHOLDER}}')

    doc.add_heading('3. Conclusion', level=1)
    doc.add_paragraph('Core task management features are complete. The remaining work focuses on '
                      'search, persistence and reporting, which are planned for the next sprint.')

    doc.save(path)
    print(f"Created {os.path.relpath(path, BASE_DIR)}")


# ---------- PDFs ----------
def draw_lines(c, lines):
    """Draws (text, style) lines; style is 'title', 'h', or 'p'."""
    width, height = A4
    y = height - 72
    fonts = {'title': ('Helvetica-Bold', 18, 28), 'h': ('Helvetica-Bold', 13, 22), 'p': ('Helvetica', 11, 16)}
    for text, style in lines:
        font, size, step = fonts[style]
        if y < 72:
            c.showPage()
            y = height - 72
        c.setFont(font, size)
        c.drawString(72, y, text)
        y -= step
    c.showPage()


def create_pdf(path, pages):
    c = canvas.Canvas(path, pagesize=A4)
    for lines in pages:
        draw_lines(c, lines)
    c.save()
    print(f"Created {os.path.relpath(path, BASE_DIR)}")


SOURCE_DOC = [[
    ("Enhanced Task Manager - Project Specification", 'title'),
    ("Document ID: ETM-SPEC-001   Version: 1.0", 'p'),
    ("", 'p'),
    ("1. Overview", 'h'),
    ("The Enhanced Task Manager is a command-line application that helps users", 'p'),
    ("organize daily work. Users can create tasks, set priorities and due dates,", 'p'),
    ("and track progress until each task is completed.", 'p'),
    ("", 'p'),
    ("2. Functional Requirements", 'h'),
    ("- Add, edit and delete tasks", 'p'),
    ("- Assign a priority (High / Medium / Low) and a due date", 'p'),
    ("- Search and filter tasks by keyword, status or priority", 'p'),
    ("- Save and load tasks from a JSON file", 'p'),
    ("- Export a summary report", 'p'),
    ("", 'p'),
    ("3. Non-Functional Requirements", 'h'),
    ("- Modular code structure that is easy to test", 'p'),
    ("- Clear error messages and logging", 'p'),
    ("- Configuration stored outside the source code", 'p'),
]]

DUMMY_1 = [[
    ("Part A - Project Planning", 'title'),
    ("Enhanced Task Manager", 'p'),
    ("", 'p'),
    ("Goals", 'h'),
    ("Build a reliable task manager with priorities, due dates and search.", 'p'),
    ("", 'p'),
    ("Milestones", 'h'),
    ("Week 1: Requirements and planning", 'p'),
    ("Week 2: Core task CRUD features", 'p'),
    ("Week 3: Search, filter and JSON persistence", 'p'),
    ("Week 4: Reporting, testing and documentation", 'p'),
]]

DUMMY_2 = [[
    ("Part B - System Design", 'title'),
    ("Enhanced Task Manager", 'p'),
    ("", 'p'),
    ("Main Components", 'h'),
    ("Task: data class holding title, priority, due date and status", 'p'),
    ("TaskManager: adds, updates, removes and searches tasks", 'p'),
    ("Storage: saves and loads tasks as JSON", 'p'),
    ("CLI: reads user commands and prints results", 'p'),
    ("", 'p'),
    ("Data Flow", 'h'),
    ("CLI -> TaskManager -> Storage (JSON file)", 'p'),
]]


if __name__ == "__main__":
    os.makedirs(DOCS_DIR, exist_ok=True)
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    create_word_template(os.path.join(TEMPLATES_DIR, 'report_template.docx'))
    create_pdf(os.path.join(DOCS_DIR, 'source_document.pdf'), SOURCE_DOC)
    create_pdf(os.path.join(DOCS_DIR, 'dummy_doc1.pdf'), DUMMY_1)
    create_pdf(os.path.join(DOCS_DIR, 'dummy_doc2.pdf'), DUMMY_2)
