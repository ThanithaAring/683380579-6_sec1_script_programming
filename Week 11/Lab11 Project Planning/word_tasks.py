# agentic-document-automator/src/word_tasks.py
import os

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT

from .utils import setup_logging

logger = setup_logging(__name__)

TABLE_PLACEHOLDER = "{{TABLE_PLACEHOLDER}}"


class WordTasks:
    """Word operations for the Document Agent: fill a template, extract text."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    @staticmethod
    def _replace_in_paragraph(paragraph, replacements):
        """
        Replaces every {{KEY}} in a paragraph.
        Word often splits text like '{{REPORT_TITLE}}' across several runs, so we
        join all runs, replace, then put the result into the first run
        (keeping that run's formatting) and clear the others.
        """
        full_text = "".join(run.text for run in paragraph.runs)
        if "{{" not in full_text:
            return
        new_text = full_text
        for key, value in replacements.items():
            new_text = new_text.replace("{{" + key + "}}", str(value))
        if new_text != full_text:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""

    @staticmethod
    def _build_table(document, table_data):
        table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
        try:
            table.style = 'Table Grid'
        except Exception:
            logger.warning("Style 'Table Grid' not in template; using default table style.")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row_data in enumerate(table_data):
            for c_idx, value in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(value)
                if r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        return table

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Fills {{KEY}} placeholders in a .docx template and inserts a table where
        {{TABLE_PLACEHOLDER}} appears (or at the end if the template has none).
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        try:
            document = docx.Document(template_path)

            # 1) Insert the dynamic table at the placeholder position
            if table_data:
                table = self._build_table(document, table_data)
                anchor = next((p for p in document.paragraphs if TABLE_PLACEHOLDER in p.text), None)
                if anchor is not None:
                    anchor._element.addnext(table._tbl)          # move table right after the anchor
                    anchor._element.getparent().remove(anchor._element)  # then drop the anchor paragraph
                    logger.info("Table inserted at {{TABLE_PLACEHOLDER}}.")
                else:
                    logger.warning("No {{TABLE_PLACEHOLDER}} in template; table appended at the end.")

            # 2) Replace text placeholders in body paragraphs and in table cells
            for paragraph in document.paragraphs:
                self._replace_in_paragraph(paragraph, data_placeholders)
            for tbl in document.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            self._replace_in_paragraph(paragraph, data_placeholders)

            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated -> {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts paragraph text and table rows (tab-separated) from a .docx."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None
        try:
            document = docx.Document(docx_path)
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
            logger.info(f"Text extracted from {os.path.basename(docx_path)}")
            return "\n".join(lines)
        except Exception as e:
            logger.error(f"Error extracting text from '{docx_path}': {e}")
            return None
